import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from actions.presentation import TOOL, presentation
from core.action_loader import _validate


def _no_open(path):
    return f"Opened '{Path(path).name}' (stub)."


def _no_powerpoint():
    return None, None


class TestToolShape(unittest.TestCase):
    def test_tool_validates(self):
        rec = _validate(TOOL, "presentation.py")
        self.assertTrue(rec.valid, rec.error)
        self.assertEqual(rec.name, "presentation")

    def test_required_action(self):
        self.assertIn("'action' is required", presentation({}))


class TestFileMode(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.path = Path(self._td.name) / "deck.pptx"
        self._patches = [
            patch("actions.presentation._open_in_powerpoint",
                  side_effect=lambda p: _no_open(p)),
            patch("actions.presentation._get_active_powerpoint",
                  side_effect=_no_powerpoint),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self._td.cleanup()

    def _count(self):
        from pptx import Presentation
        return len(Presentation(str(self.path)).slides._sldIdLst)

    def test_create_two_slides_with_content(self):
        out = presentation({
            "action": "create",
            "title": "My Deck",
            "path": str(self.path),
            "slides": [
                {"title": "Page 1", "bullets": ["Hello world"]},
                {"title": "Page 2", "bullets": ["Line A", "Line B"]},
            ],
        })
        self.assertIn("Created 2-slide deck", out)
        self.assertTrue(self.path.exists())
        self.assertEqual(self._count(), 2)

        from pptx import Presentation
        prs = Presentation(str(self.path))
        self.assertEqual(prs.slides[0].shapes.title.text, "Page 1")
        self.assertEqual(prs.slides[1].shapes.title.text, "Page 2")
        body = prs.slides[1].placeholders[1].text
        self.assertIn("Line A", body)
        self.assertIn("Line B", body)

    def test_create_opens_the_file(self):
        out = presentation({
            "action": "create", "title": "T", "path": str(self.path),
            "slides": [{"title": "Only", "bullets": []}],
        })
        self.assertIn("deck.pptx", out)      # path reported
        self.assertIn("stub", out)           # open step actually ran

    def test_create_defaults_to_one_titled_slide(self):
        out = presentation({"action": "create", "title": "Solo",
                            "path": str(self.path)})
        self.assertIn("Created 1-slide deck", out)
        self.assertEqual(self._count(), 1)

    def test_add_slide_appends(self):
        presentation({"action": "create", "title": "T", "path": str(self.path),
                      "slides": [{"title": "One", "bullets": []}]})
        out = presentation({"action": "add_slide", "title": "Two",
                            "bullets": ["point"], "path": str(self.path)})
        self.assertIn("Added slide 2", out)
        self.assertEqual(self._count(), 2)

    def test_write_edits_existing_slide(self):
        presentation({"action": "create", "title": "T", "path": str(self.path),
                      "slides": [{"title": "One", "bullets": ["old"]},
                                 {"title": "Two", "bullets": []}]})
        out = presentation({"action": "write", "slide": 1,
                            "title": "Renamed", "text": "fresh",
                            "path": str(self.path)})
        self.assertIn("Wrote slide 1", out)
        from pptx import Presentation
        prs = Presentation(str(self.path))
        self.assertEqual(prs.slides[0].shapes.title.text, "Renamed")
        self.assertIn("fresh", prs.slides[0].placeholders[1].text)

    def test_write_slide_out_of_range(self):
        presentation({"action": "create", "title": "T", "path": str(self.path),
                      "slides": [{"title": "One", "bullets": []}]})
        out = presentation({"action": "write", "slide": 9, "text": "x",
                            "path": str(self.path)})
        self.assertIn("does not exist", out)

    def test_status_reports_file(self):
        presentation({"action": "create", "title": "T", "path": str(self.path),
                      "slides": [{"title": "A", "bullets": []},
                                 {"title": "B", "bullets": []}]})
        out = presentation({"action": "status", "path": str(self.path)})
        self.assertIn("2 slides", out)
        self.assertIn("not running", out)

    def test_unknown_action(self):
        out = presentation({"action": "shuffle", "path": str(self.path)})
        self.assertIn("Unknown action", out)

    def test_open_without_file(self):
        with patch("actions.presentation._newest_our_pptx",
                   return_value=None):
            out = presentation({"action": "open"})
        self.assertIn("No presentation file", out)


if __name__ == "__main__":
    unittest.main()
